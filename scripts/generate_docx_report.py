import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Base Colors
    PRIMARY = RGBColor(16, 185, 129)    # Emerald 500
    DARK_NAVY = RGBColor(15, 23, 42)    # Slate 900
    TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
    TEXT_BODY = RGBColor(30, 41, 59)    # Slate 800

    # ------------------ COVER HEADER ------------------
    p_pre = doc.add_paragraph()
    r_pre = p_pre.add_run("ENTERPRISE COMPLIANCE ARCHITECTURE · MVP RELEASE v1.0")
    r_pre.font.name = "Arial"
    r_pre.font.size = Pt(10)
    r_pre.font.bold = True
    r_pre.font.color.rgb = PRIMARY

    title_p = doc.add_paragraph()
    title_r = title_p.add_run("Carbon Crunch — Compliance Management Platform")
    title_r.font.name = "Arial"
    title_r.font.size = Pt(24)
    title_r.font.bold = True
    title_r.font.color.rgb = DARK_NAVY

    sub_p = doc.add_paragraph()
    sub_r = sub_p.add_run("Comprehensive Handover Documentation, Role-Based Access Control Guide & Technical Specification")
    sub_r.font.name = "Arial"
    sub_r.font.size = Pt(12)
    sub_r.font.color.rgb = TEXT_MUTED

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ------------------ SECTION 1: CREDENTIALS MATRIX ------------------
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("1. Authenticated User Accounts & Access Credentials")
    h1_run.font.name = "Arial"
    h1_run.font.size = Pt(16)
    h1_run.font.bold = True
    h1_run.font.color.rgb = DARK_NAVY

    desc_p = doc.add_paragraph()
    desc_p.add_run("The platform enforces strict Role-Based Access Control (RBAC) backed by HTTP-only secure cookie JWT sessions (`cc_auth_token`). Below is the verified matrix of pre-seeded credentials for executive testing:")
    desc_p.paragraph_format.space_after = Pt(8)

    # Credentials Table
    headers = ["Role", "User Name", "Email Address", "Password", "Designation / Unit", "Access Rights & Authorizations"]
    data = [
        ["CONSULTANT", "Priya Sharma", "consultant@carboncrunch.com", "Demo@123", "Senior Energy Consultant", "Multi-facility matrices, Excel reconciliation, DISCOM tolerance check, email generator. Cannot approve datasets."],
        ["REVIEWER", "Dr. Rajesh Nair", "reviewer@carboncrunch.com", "Demo@123", "Chief Compliance Auditor", "Reviewer Queue (/approvals). Authorized to independently audit, sign off, lock versions, or reject with notes."],
        ["CLIENT", "Vikram Singh", "client@tridentindia.com", "Demo@123", "Head of Regulatory, Trident", "Confined to Client Portal (/portal). Can inspect Trident facilities, submit Google Drive links, answer queries."],
        ["ADMIN", "Saurabh Agnihotri", "admin@carboncrunch.com", "Demo@123", "Principal Compliance Architect", "Universal platform access, conglomerate onboarding, system logs, database reseeding, user management."]
    ]

    table = doc.add_table(rows=len(data) + 1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(1.1), Inches(1.0), Inches(1.8), Inches(0.8), Inches(1.3), Inches(1.8)]

    # Header Row
    for col_idx, text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.width = col_widths[col_idx]
        set_cell_background(cell, "0F172A") # Slate 900
        set_cell_margins(cell, 120, 120, 120, 120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for row_idx, row_data in enumerate(data):
        bg_hex = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.width = col_widths[col_idx]
            set_cell_background(cell, bg_hex)
            set_cell_margins(cell, 100, 100, 100, 100)
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            if col_idx == 0:
                run.font.bold = True
                run.font.color.rgb = PRIMARY
            elif col_idx == 3:
                run.font.bold = True
                run.font.color.rgb = RGBColor(225, 29, 72)
            else:
                run.font.color.rgb = TEXT_BODY

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # ------------------ SECTION 2: SECURITY & RBAC ARCHITECTURE ------------------
    h2 = doc.add_heading(level=1)
    h2_run = h2.add_run("2. Multi-Tenant Security & RBAC Enforcement")
    h2_run.font.name = "Arial"
    h2_run.font.size = Pt(16)
    h2_run.font.bold = True
    h2_run.font.color.rgb = DARK_NAVY

    p2 = doc.add_paragraph()
    p2.add_run("1. JWT Token Cookies: ").bold = True
    p2.add_run("Upon authentication, the server writes a cryptographically signed JWT to an HTTP-only secure cookie (`cc_auth_token`). When a user logs out, the cookie is immediately evicted by `/api/auth/logout`.")

    p3 = doc.add_paragraph()
    p3.add_run("2. Next.js Security Middleware Guard (src/middleware.ts): ").bold = True
    p3.add_run("Unauthorized visitors are automatically redirected to `/login`. Furthermore, role boundaries are strictly enforced:\n")
    p3.add_run("   • External CLIENT users are strictly confined to `/portal`. Direct URL navigation to `/dashboard`, `/approvals`, `/facilities/...`, or `/datasets/...` is intercepted and blocked.\n")
    p3.add_run("   • CONSULTANT users cannot access the independent Reviewer Queue (`/approvals`). Under compliance integrity guidelines, consultants cannot approve their own submissions.\n")
    p3.add_run("   • Internal CONSULTANT and REVIEWER users cannot enter the external Client Portal (`/portal`).")

    p4 = doc.add_paragraph()
    p4.add_run("3. UI Role Isolation: ").bold = True
    p4.add_run("Arbitrary role-switching buttons in the header have been completely removed. The navigation menu renders only routes explicitly permitted for the authenticated role. To change roles, users must click Logout and log in as that designated account.")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ------------------ SECTION 3: CORE COMPLIANCE MODULES ------------------
    h3 = doc.add_heading(level=1)
    h3_run = h3.add_run("3. Core Functional Modules (MVP Specification)")
    h3_run.font.name = "Arial"
    h3_run.font.size = Pt(16)
    h3_run.font.bold = True
    h3_run.font.color.rgb = DARK_NAVY

    modules = [
        ("Module 1: Multi-Facility Compliance Matrices", 
         "Centralizes RCO/RPO, PAT, CCTS, and CEA returns across industrial units (Trident Paper, Yarn, Textiles). Tracks status (Compliant, Shortfall, In Progress), quantifies shortfalls (5.0 MU RCO FY24-25), and plots YoY trajectory tables."),
        ("Module 2: Google Drive Evidence Integration", 
         "Compliance values are linked to external Google Drive documents. Clients upload Drive share links for monthly meter logs, DISCOM bills, and calibration records. Carbon Crunch tracks metadata without hosting files on local disk."),
        ("Module 3: Verification & Reconciliation Engine", 
         "Reconciles client-declared values against DISCOM utility bills with automated tolerance checks. In June 2025 demo case, flags a +15,000 kWh (1.186%) mismatch between declared 1,280,000 kWh and billed 1,265,000 kWh."),
        ("Module 4: Clarification & Query Lifecycle", 
         "Includes a 1-click official discrepancy email generator pre-populated with consumer IDs, meter numbers, and variance percentages. Clients receive portal alerts and submit explanations with revised Drive evidence."),
        ("Module 5: Independent Reviewer Queue & Audit Locking", 
         "Chief Reviewer Dr. Rajesh Nair reviews consultant reconciliations. Upon sign-off, the dataset version is locked against further editing, stamping reviewer identity, comments, and audit timestamp."),
        ("Module 6: Statutory Compliance Calendar", 
         "Tracks regulatory due dates, client cut-off milestones, and CEA recurring monthly returns. Features deadline extension tracking with circular reference numbers, regulatory reasons, and extension counters."),
        ("Module 7: Cross-Role Notification Engine", 
         "Dispatches real-time triggers between Consultant, Reviewer, and Client when actions occur (clarification raised, response submitted, review requested, audit approved). Features an interactive Notification Bell with unread counter in the top bar.")
    ]

    for title, desc in modules:
        mp = doc.add_paragraph()
        mr = mp.add_run(f"• {title}: ")
        mr.bold = True
        mr.font.color.rgb = PRIMARY
        mp.add_run(desc)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ------------------ SECTION 4: IN-PLATFORM DOCUMENT VIEWER ------------------
    h4 = doc.add_heading(level=1)
    h4_run = h4.add_run("4. Evidence Document Previewer & Google Drive Handling")
    h4_run.font.name = "Arial"
    h4_run.font.size = Pt(16)
    h4_run.font.bold = True
    h4_run.font.color.rgb = DARK_NAVY

    p_doc = doc.add_paragraph()
    p_doc.add_run("Why Google Drive 404 occurs in mock links: ").bold = True
    p_doc.add_run("Because seeded demo URLs point to placeholder file paths (e.g., `PSPCL_May2025_Bill.pdf`), navigating to `drive.google.com` prompted Google's real servers to return a 'File does not exist' 404 page.\n\n")
    p_doc.add_run("Solution Implemented: ").bold = True
    p_doc.add_run("Clicking any 'PSPCL Bill' badge launches the in-platform Evidence Document Modal, rendering a high-fidelity official Punjab State Power Corporation Limited (PSPCL) bill with Account No. `PSPCL-BARNALA-66KV-0042`, Meter No. `L&T-EM-99214`, sanctioned demand (12,000 kVA @ 66 kV), metered units, and discrepancy analysis, alongside a direct button to open the link in Drive.")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ------------------ SECTION 5: TECHNICAL STACK & API REFERENCE ------------------
    h5 = doc.add_heading(level=1)
    h5_run = h5.add_run("5. Technical Architecture & REST API Endpoints")
    h5_run.font.name = "Arial"
    h5_run.font.size = Pt(16)
    h5_run.font.bold = True
    h5_run.font.color.rgb = DARK_NAVY

    api_headers = ["Endpoint", "HTTP Method", "Authorized Roles", "Functional Description"]
    api_data = [
        ["/api/auth/login", "POST", "Public", "Validates user credentials, generates signed JWT, sets HTTP-only cookie."],
        ["/api/auth/logout", "POST", "Authenticated", "Evicts cc_auth_token cookie, terminates authenticated session."],
        ["/api/auth/me", "GET", "Authenticated", "Returns session details (userId, name, role, clientGroupId)."],
        ["/api/seed", "POST", "Admin", "Reseeds MongoDB with complete Trident Group test records and notifications."],
        ["/api/facilities/[id]/compliance-matrix", "GET", "Internal / Client", "Fetches multi-year compliance status, shortfalls, and matrix grid."],
        ["/api/datasets", "GET", "Internal", "Retrieves datasets for reconciliation and reviewer queues."],
        ["/api/datasets/[id]", "GET, PATCH", "Internal", "Fetches or updates monthly meter consumption and evidence links."],
        ["/api/verifications/[id]/approve", "POST", "Reviewer / Admin", "Signs off dataset, locks version, and triggers approval notifications."],
        ["/api/verifications/[id]/reject", "POST", "Reviewer / Admin", "Rejects dataset with audit notes and notifies consultant."],
        ["/api/clarifications", "GET, POST", "Internal / Client", "Lists or initiates discrepancy clarification queries."],
        ["/api/clarifications/[id]", "PATCH", "Client / Internal", "Records client explanation and revised Drive link."],
        ["/api/calendar", "GET, PATCH", "Internal", "Returns compliance calendar or logs statutory deadline extensions."],
        ["/api/notifications", "GET, POST", "Authenticated", "Retrieves unread notifications or creates cross-role triggers."],
        ["/api/notifications/mark-all", "PATCH", "Authenticated", "Marks all notifications as read for current active role."],
        ["/api/portal/dashboard", "GET", "Client / Admin", "Returns client group facilities, outstanding requests, and history."]
    ]

    api_table = doc.add_table(rows=len(api_data) + 1, cols=4)
    api_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    api_table.autofit = False

    api_col_widths = [Inches(2.5), Inches(1.1), Inches(1.4), Inches(2.8)]

    for col_idx, text in enumerate(api_headers):
        cell = api_table.cell(0, col_idx)
        cell.width = api_col_widths[col_idx]
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 120, 120, 120, 120)
        p = cell.paragraphs[0]
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    for row_idx, row_data in enumerate(api_data):
        bg_hex = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            cell = api_table.cell(row_idx + 1, col_idx)
            cell.width = api_col_widths[col_idx]
            set_cell_background(cell, bg_hex)
            set_cell_margins(cell, 90, 90, 90, 90)
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            if col_idx == 0:
                run.font.bold = True
                run.font.color.rgb = DARK_NAVY
            elif col_idx == 1:
                run.font.bold = True
                run.font.color.rgb = PRIMARY
            else:
                run.font.color.rgb = TEXT_BODY

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # ------------------ SECTION 6: UI DESIGN SYSTEM & DUAL THEMES ------------------
    h6 = doc.add_heading(level=1)
    h6_run = h6.add_run("6. User Interface Design System & Dual Theme Architecture")
    h6_run.font.name = "Arial"
    h6_run.font.size = Pt(16)
    h6_run.font.bold = True
    h6_run.font.color.rgb = DARK_NAVY

    p_theme1 = doc.add_paragraph()
    p_theme1.add_run("1. High-Contrast Light Mode: ").bold = True
    p_theme1.add_run("Features a dedicated, crisp white navbar (`bg-white/95 border-b border-slate-200`) with dark slate typography (`text-slate-900`, `text-slate-700`), a clean `bg-slate-100` status bar, high-contrast table headers (`thead` in `bg-slate-100`), vibrant red discrepancy action cards, and readable badges with emerald and rose accents.")

    p_theme2 = doc.add_paragraph()
    p_theme2.add_run("2. Untouched Dark Mode: ").bold = True
    p_theme2.add_run("Completely preserves the dark navy slate (`#020617`, `#0f172a`) aesthetic with glowing emerald highlights, subtle borders, and dark glassmorphic panels without any visual regressions.")

    p_theme3 = doc.add_paragraph()
    p_theme3.add_run("3. Password Visibility Toggle: ").bold = True
    p_theme3.add_run("The login interface includes an interactive eye icon button allowing users to toggle between masked (`••••••••`) and unmasked plaintext password input.")

    p_theme4 = doc.add_paragraph()
    p_theme4.add_run("4. Source Code Management (.gitignore): ").bold = True
    p_theme4.add_run("Standardized `.gitignore` excludes `.next/` build outputs, `node_modules/`, local `.env` secrets, TypeScript build artifacts, and debug logs from version control repositories.")

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    # Footer note
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("Carbon Crunch Compliance Platform MVP v1.0 · Confidential Handover Report")
    fr.font.name = "Arial"
    fr.font.size = Pt(8.5)
    fr.font.color.rgb = TEXT_MUTED

    # Save to disk
    output_path = r"c:\Users\saura\Desktop\new CMS CC\Carbon_Crunch_Compliance_Platform_v1.0_Report.docx"
    doc.save(output_path)
    print(f"Successfully generated Word document at: {output_path}")

if __name__ == "__main__":
    create_report()
